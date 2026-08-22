"""Excel export: one sheet per pathology, bilingual headers."""

from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook

from src import schema
from src.records import list_records


def export_xlsx(pathology_key: str | None = None) -> bytes:
    keys = [pathology_key] if pathology_key else list(schema.PATHOLOGIES)
    for k in keys:
        if k not in schema.PATHOLOGIES:
            raise KeyError(f"unknown pathology '{k}'")

    wb = Workbook()
    wb.remove(wb.active)
    for key in keys:
        spec = schema.PATHOLOGIES[key]
        ws = wb.create_sheet(spec["nombre_es"])
        fields = spec["fields"]

        ws.append(["_id", "patología / pathology"]
                  + [f"{f['label_es']} / {f['label_en']}" for f in fields])
        records_list, _total = list_records(pathology_key=key, limit=100_000)
        for r in records_list:
            ws.append([r["_id"], key] + [_cell(r["fields"][f["id"]]) for f in fields])

        for col_idx, f in enumerate(fields, start=3):
            width = max(len(f["label_es"]), len(f["label_en"]), 12)
            letter = ws.cell(row=1, column=col_idx).column_letter
            ws.column_dimensions[letter].width = min(width + 2, 40)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell(value: Any) -> Any:
    return "" if value is None else value
